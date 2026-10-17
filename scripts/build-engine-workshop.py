"""Builds the downloads of workshop 8.28 into public/downloads/:

    engine-workshop-template.xlsx   the small engine in a spreadsheet, for the learner to build (row 1 of the search grid
                                    is the worked example; the rest of the formula columns are empty)
    engine-workshop-solution.xlsx   every formula filled
    engine-workshop.ipynb           the same engine in Python, cell by cell, explained line by line

    python3 scripts/build-engine-workshop.py            # writes the three files
    python3 scripts/build-engine-workshop.py --verify   # also recalculates the sheets in LibreOffice and runs the notebook's code

Nothing here types a number of the case or a result. The case is the fact `ws.case` of data/lessons/8-28.json (written by
scripts/lessons/lesson-8-28.ts from src/data/case-s8.ts); the expected results are that lesson's facts (src/lib/calc/ws8.ts);
the notebook's code is checks/ws8_engine.py, the same file the independent check runs. Colours come from src/styles/tokens.css.
The output is reproducible: fixed document dates and zip timestamps. The sheet is an APPROXIMATION of the engine (exact
payments, no rounding of each month's interest): --verify measures how far, and the sheet's own check tab shows it.

    Sheet "מנוע": the case (inputs), your choices, the banks, the rules, the positioning, the results and the explanation,
                  then the search grid (one row per candidate)
    Sheet "בדיקה": the sheet against the course's engine, for the default inputs
    Sheet "הוראות": what to write
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "checks"))
OUT = ROOT / "public" / "downloads"
DATA = ROOT / "data" / "lessons" / "8-28.json"
TOKENS = ROOT / "src" / "styles" / "tokens.css"
FIXED_DATE = "2026-10-17T00:00:00Z"
ZIP_TIME = (2026, 10, 17, 0, 0, 0)
MARK = "דוגמה לימודית — נתוני הלקוח והריביות להמחשה בלבד; אינם הצעת בנק או נתוני שוק עדכניים. נתונים סינתטיים. עתיד: זה לא קוד של Mogi."


def facts():
    return {f["id"]: f for f in json.loads(DATA.read_text(encoding="utf-8"))["facts"]}


def tokens():
    css = TOKENS.read_text(encoding="utf-8")
    return {m.group(1): m.group(2).upper() for m in re.finditer(r"--mogi-([a-z0-9-]+):\s*#([0-9A-Fa-f]{6})", css)}


IDS = ["sum_equals_loan", "equity_available", "ltv", "pti", "variable_share", "term", "bank_ratio", "bank_ltv", "client_max_payment", "client_target_ratio"]
IDS_HE = {
    "sum_equals_loan": "סכום המסלולים שווה להלוואה",
    "equity_available": "ההון העצמי קיים",
    "ltv": "שיעור מימון עד התקרה",
    "pti": "יחס החזר עד 50%",
    "variable_share": "לפחות שליש בריבית קבועה",
    "term": "עד 30 שנה",
    "bank_ratio": "יחס ההחזר עד הקו של הבנק",
    "bank_ltv": "שיעור המימון עד הקו של הבנק",
    "client_max_payment": "ההחזר עד המרבי שהלקוח קבע",
    "client_target_ratio": "יחס ההחזר עד יעד הלקוח",
}


def num(x):
    d = Decimal(str(x))
    return int(d) if d == d.to_integral() else float(d)


# ── the layout ─────────────────────────────────────────────────────────────


def build(kind, F):
    case = F["ws.case"]["value"]
    rules = case["rules"]
    t = tokens()
    navy, white = t["navy"], t["n0"]
    font = lambda **kw: Font(name="Arial", size=kw.pop("size", 11), **kw)  # noqa: E731
    fill = lambda hex_: PatternFill("solid", fgColor=hex_)  # noqa: E731
    thin = Side(style="thin", color=t["n200"])
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    money, pct, plain = "#,##0.00", "0.00", "0"

    wb = Workbook()
    ws = wb.active
    ws.title = "מנוע"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 38
    for i in range(2, 42):
        ws.column_dimensions[get_column_letter(i)].width = 15

    solution = kind == "solution"
    ws["A1"] = "מנוע קטן לבחירת תמהיל — סדנה 8.28" + (" (פתור)" if solution else "")
    ws["A1"].font = font(size=14, bold=True, color=navy)
    ws["A2"] = MARK
    ws["A2"].font = font(size=10, color=t["n600"])

    A = {}  # name -> absolute address like $B$5

    def header(row, text):
        ws[f"A{row}"] = text
        ws[f"A{row}"].font = font(bold=True, color=navy)

    def put(row, name, label, value, fmt=None, inp=False, note=None):
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = font(color=t["n800"])
        c = ws[f"B{row}"]
        c.value = value
        if fmt:
            c.number_format = fmt
        if inp:
            c.fill = fill(t["mint"])
            c.border = box
            c.font = font(bold=True, color=navy)
        if note:
            ws[f"C{row}"] = note
            ws[f"C{row}"].font = font(size=10, color=t["n600"])
        A[name] = f"$B${row}"

    # inputs: the case
    header(4, "נתוני התיק (סינתטיים, מתוך התיק של השלב)")
    put(5, "price", "מחיר הנכס (₪)", num(case["price"]), money, True)
    put(6, "base_eq", "ההון העצמי של הבסיס (₪)", num(case["base_equity"]), money, True)
    put(7, "savings", "חסכונות (₪)", num(case["savings"]), money, True)
    put(8, "payoff", "תשלום לסגירת הלוואת הרכב (₪)", num(case["car_payoff"]), money, True)
    put(9, "car", "החזר חודשי של הלוואת הרכב (₪)", num(case["car_monthly"]), money, True)
    put(10, "car_left", "חודשים שנשארו להלוואת הרכב", case["car_months_left"], plain, True)
    put(11, "inc1", "הכנסה נטו של לווה 1 (₪)", num(case["incomes"][0]), money, True)
    put(12, "inc2", "הכנסה נטו של לווה 2 (₪)", num(case["incomes"][1]), money, True)
    put(13, "parent_inc", "הכנסת הורה, אם מצטרף (₪)", num(case["parent_income"]), money, True)
    put(14, "maxpay", "ההחזר המרבי שהלקוח קבע (₪)", num(case["max_payment"]), money, True)
    put(15, "target", "יעד יחס ההחזר של הלקוח (%)", num(case["target_ratio"]), pct, True)

    header(17, "הבחירות שלכם — מיצוב הבקשה (שנו כאן)")
    put(18, "extra", "הון עצמי נוסף מהחסכונות (₪)", num(json_choice(F, "extra_equity")), money, True, "הערך ההתחלתי הוא הבחירה של שיעור 8.28")
    put(19, "close", "סוגרים את הלוואת הרכב? (1 כן, 0 לא)", int(json_choice(F, "close_car")), plain, True)
    put(20, "parent", "הורה מצטרף כלווה? (1 כן, 0 לא)", int(json_choice(F, "parent")), plain, True)
    put(21, "bank", "בנק (a, b, c או d)", json_choice(F, "bank"), None, True, "ארבעה בנקים בדויים")
    weights = F["p.winner_id"]["inputs"]["weights"].split("|")
    put(22, "w1", "משקל: ההחזר הראשון", int(weights[0]), plain, True, "להדגמה: לא החלטה של Mogi")
    put(23, "w2", "משקל: סך התשלומים", int(weights[1]), plain, True)
    put(24, "w3", "משקל: החשיפה", int(weights[2]), plain, True)

    header(26, "הבנקים הבדויים: מחירון (אחוז לשנה) וקווי סירוב")
    for col, text in zip("ABCDEFG", ("קוד", "שם", "קבועה", "פריים", "משתנה", "סירוב ביחס החזר (%)", "סירוב בשיעור מימון (%)")):
        ws[f"{col}27"] = text
        ws[f"{col}27"].font = font(bold=True, color=white)
        ws[f"{col}27"].fill = fill(navy)
        ws[f"{col}27"].alignment = Alignment(horizontal="center", wrap_text=True)
    for k, (bid, b) in enumerate(case["banks"].items()):
        r = 28 + k
        vals = (bid, b["name"], num(b["prices"]["fixed_unlinked"]), num(b["prices"]["prime"]), num(b["prices"]["variable_unlinked"]), num(b["ratio_max"]), num(b["ltv_max"]))
        for col, v in zip("ABCDEFG", vals):
            ws[f"{col}{r}"] = v
            ws[f"{col}{r}"].border = box
            ws[f"{col}{r}"].font = font(color=t["n800"])
    bank_rows = "$A$28:$A$31"

    header(33, "הכללים (כפי שתועדו בקורס; הקווים של הרגולציה: לא אומת)")
    put(34, "ltv_ceil", "תקרת שיעור מימון (%)", num(rules["ltv_ceiling"]), pct)
    put(35, "pti_lim", "תקרת יחס ההחזר (%)", num(rules["pti_limit"]), pct)
    put(36, "var_num", "החלק המרבי בריבית משתנה: מונה", rules["variable_numerator"], plain)
    put(37, "var_den", "החלק המרבי בריבית משתנה: מכנה", rules["variable_denominator"], plain)
    put(38, "max_term", "התקופה המרבית (חודשים)", rules["max_term_months"], plain)
    put(39, "n", "התקופה של כל מסלול בחיפוש (חודשים)", rules["years"] * 12, plain)
    put(40, "s_prime", "תרחיש לחץ: פריים עולה (נקודות אחוז, מהשנה השנייה)", num(case["scenarios"]["stress"]["prime_shift"]), pct)
    put(41, "s_anchor", "תרחיש לחץ: העוגן עולה בעדכון (נקודות אחוז)", num(case["scenarios"]["stress"]["anchor_shift"]), pct)
    put(42, "reset", "חודשים בין עדכוני הריבית במסלול משתנה", rules["reset_months"], plain)

    header(44, "מיצוב: מהבחירות לבקשה")
    put(45, "equity", "ההון העצמי שנכנס לעסקה (₪)", f"={A['base_eq']}+{A['extra']}", money)
    put(46, "avail", "ההון העצמי הזמין (₪)", f"={A['base_eq']}+MAX({A['savings']}-IF({A['close']}=1,{A['payoff']},0),0)", money)
    put(47, "loan", "סכום ההלוואה (₪)", f"={A['price']}-{A['equity']}", money)
    put(48, "ltv", "שיעור מימון (%)", f"={A['loan']}/{A['price']}*100", pct)
    put(49, "income", "הכנסה מוכרת (₪)", f"={A['inc1']}+{A['inc2']}+IF({A['parent']}=1,{A['parent_inc']},0)", money)
    put(50, "car_counts", "הלוואת הרכב נספרת בחישוב (₪, כשנשארו מעל 18 חודשים)", f"=IF(AND({A['close']}=0,{A['car_left']}>18),{A['car']},0)", money)
    put(51, "disp", "הכנסה פנויה (₪)", f"={A['income']}-{A['car_counts']}", money)
    brow = f"MATCH({A['bank']},{bank_rows},0)"
    put(52, "rf", "מחירון הבנק: קבועה (%)", f"=INDEX($C$28:$C$31,{brow})", pct)
    put(53, "rp", "מחירון הבנק: פריים (%)", f"=INDEX($D$28:$D$31,{brow})", pct)
    put(54, "rv", "מחירון הבנק: משתנה (%)", f"=INDEX($E$28:$E$31,{brow})", pct)
    put(55, "b_ratio", "קו הסירוב של הבנק ביחס החזר (%)", f"=INDEX($F$28:$F$31,{brow})", pct)
    put(56, "b_ltv", "קו הסירוב של הבנק בשיעור מימון (%)", f"=INDEX($G$28:$G$31,{brow})", pct)

    # the grid
    shares = [tuple(int(x) for x in re.findall(r"\d+", gid)) for gid in grid_ids(F)]
    HEAD = 88
    FIRST = HEAD + 1
    LAST = HEAD + len(shares)
    cols = [
        ("id", "מועמד", None), ("f", "קבועה %", None), ("p", "פריים %", None), ("v", "משתנה %", None),
        ("ef", "סכום קבועה", money), ("ep", "סכום פריים", money), ("ev", "סכום משתנה", money),
        ("hf", "החזר ראשון קבועה", money), ("hp", "החזר ראשון פריים", money), ("hv", "החזר ראשון משתנה", money),
        ("first", "ההחזר הראשון", money), ("qual", "ההחזר לחישוב היחס (בלי עיגול)", money), ("total", "סך התשלומים", money),
        ("stress", "ההחזר הגבוה בלחץ", money), ("expo", "חשיפה", money), ("varshare", "חלק משתנה %", pct), ("ratio", "יחס החזר %", pct),
    ] + [(f"x_{i}", IDS_HE[i], plain) for i in IDS] + [
        ("nrs", "מספר סיבות פסילה", plain), ("feas", "אפשרי? (1 כן)", plain), ("dom", "כמה אפשריים שולטים עליו", plain), ("eff", "יעיל? (1 כן)", plain),
        ("o1", "ההחזר הראשון (אפשריים)", money), ("o2", "סך התשלומים (אפשריים)", money), ("o3", "חשיפה (אפשריים)", money),
        ("n1", "מנורמל: ראשון", "0.0000"), ("n2", "מנורמל: סך", "0.0000"), ("n3", "מנורמל: חשיפה", "0.0000"),
        ("score", "ציון משוקלל", "0.000000"), ("rank", "דירוג בין האפשריים", plain),
    ]
    L = {key: get_column_letter(i + 1) for i, (key, _, _) in enumerate(cols)}
    rng = lambda key: f"${L[key]}${FIRST}:${L[key]}${LAST}"  # noqa: E731

    header(HEAD - 2, "חיפוש: רשת של תמהילים (כל חלק כפולה של 10%, כל המסלולים 30 שנה), אילוצים, חזית ובחירה")
    for i, (key, label, _) in enumerate(cols):
        c = ws.cell(row=HEAD, column=i + 1, value=label)
        c.font = font(bold=True, color=white)
        c.fill = fill(navy)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[HEAD].height = 48

    # normalisation bounds, above the grid
    lo_hi = {}
    for j, (src, onl) in enumerate((("first", "o1"), ("total", "o2"), ("expo", "o3"))):
        r = HEAD - 5 + j
        ws[f"A{r}"] = ("הטוב ביותר ", "הגרוע ביותר ")[0] + ("ההחזר הראשון", "סך התשלומים", "החשיפה")[j] + " (בין האפשריים)"
        lo = f"=IF(COUNT({rng(onl)})=0,MIN({rng(src)}),MIN({rng(onl)}))"
        hi = f"=IF(COUNT({rng(onl)})=0,MAX({rng(src)}),MAX({rng(onl)}))"
        ws[f"B{r}"], ws[f"C{r}"] = lo, hi
        for col in "BC":
            ws[f"{col}{r}"].number_format = money
        lo_hi[j] = (f"$B${r}", f"$C${r}")
    ws[f"B{HEAD - 6}"], ws[f"C{HEAD - 6}"] = "הנמוך", "הגבוה"

    def formulas(r):
        n = A["n"]
        pm = lambda rate, amt: f"IF({amt}>0,-PMT({rate}/1200,{n},{amt}),0)"  # noqa: E731
        bal = lambda amt, rate, k: f"{amt}*((1+{rate}/1200)^{n}-(1+{rate}/1200)^{k})/((1+{rate}/1200)^{n}-1)"  # noqa: E731
        c = lambda key: f"{L[key]}{r}"  # noqa: E731
        f = {
            "ef": f"=ROUND({A['loan']}*{c('f')}/100,2)",
            "ep": f"=ROUND({A['loan']}*{c('p')}/100,2)",
            "ev": f"=ROUND({A['loan']}*{c('v')}/100,2)",
            "hf": f"=IF({c('ef')}>0,ROUND(-PMT({A['rf']}/1200,{n},{c('ef')}),2),0)",
            "hp": f"=IF({c('ep')}>0,ROUND(-PMT({A['rp']}/1200,{n},{c('ep')}),2),0)",
            "hv": f"=IF({c('ev')}>0,ROUND(-PMT({A['rv']}/1200,{n},{c('ev')}),2),0)",
            "first": f"={c('hf')}+{c('hp')}+{c('hv')}",
            "qual": f"={pm(A['rf'], c('ef'))}+{pm(A['rp'], c('ep'))}+{pm(A['rv'], c('ev'))}",
            "total": f"={c('qual')}*{n}",
            "stress": (
                f"={pm(A['rf'], c('ef'))}"
                f"+IF({c('ep')}>0,-PMT(({A['rp']}+{A['s_prime']})/1200,{n}-12,{bal(c('ep'), A['rp'], 12)}),0)"
                f"+IF({c('ev')}>0,-PMT(({A['rv']}+{A['s_anchor']})/1200,{n}-{A['reset']},{bal(c('ev'), A['rv'], A['reset'])}),0)"
            ),
            "expo": f"={c('stress')}-{c('first')}",
            "varshare": f"={c('p')}+{c('v')}",
            "ratio": f"=IF({A['disp']}>0,{c('qual')}/{A['disp']}*100,1000000)",
            "x_sum_equals_loan": f"=IF(ABS({c('ef')}+{c('ep')}+{c('ev')}-{A['loan']})>0.005,1,0)",
            "x_equity_available": f"=IF({A['equity']}>{A['avail']},1,0)",
            "x_ltv": f"=IF({A['ltv']}>{A['ltv_ceil']},1,0)",
            "x_pti": f"=IF({c('ratio')}>{A['pti_lim']},1,0)",
            "x_variable_share": f"=IF(({c('ep')}+{c('ev')})*{A['var_den']}>{A['loan']}*{A['var_num']},1,0)",
            "x_term": f"=IF({n}>{A['max_term']},1,0)",
            "x_bank_ratio": f"=IF({c('ratio')}>{A['b_ratio']},1,0)",
            "x_bank_ltv": f"=IF({A['ltv']}>{A['b_ltv']},1,0)",
            "x_client_max_payment": f"=IF({c('qual')}>{A['maxpay']},1,0)",
            "x_client_target_ratio": f"=IF({c('ratio')}>{A['target']},1,0)",
            "nrs": f"=SUM({L['x_sum_equals_loan']}{r}:{L['x_client_target_ratio']}{r})",
            "feas": f"=IF({c('nrs')}=0,1,0)",
            "dom": (
                f'=IF({c("feas")}=1,SUMPRODUCT(({rng("feas")}=1)*({rng("first")}<={c("first")})*({rng("total")}<={c("total")})*({rng("expo")}<={c("expo")})'
                f'*((({rng("first")}<{c("first")})+({rng("total")}<{c("total")})+({rng("expo")}<{c("expo")}))>0)),"")'
            ),
            "eff": f"=IF({c('feas')}=1,IF({c('dom')}=0,1,0),0)",
            "o1": f'=IF({c("feas")}=1,{c("first")},"")',
            "o2": f'=IF({c("feas")}=1,{c("total")},"")',
            "o3": f'=IF({c("feas")}=1,{c("expo")},"")',
            "n1": f"=IF({lo_hi[0][1]}={lo_hi[0][0]},0,ROUND(({c('first')}-{lo_hi[0][0]})/({lo_hi[0][1]}-{lo_hi[0][0]}),4))",
            "n2": f"=IF({lo_hi[1][1]}={lo_hi[1][0]},0,ROUND(({c('total')}-{lo_hi[1][0]})/({lo_hi[1][1]}-{lo_hi[1][0]}),4))",
            "n3": f"=IF({lo_hi[2][1]}={lo_hi[2][0]},0,ROUND(({c('expo')}-{lo_hi[2][0]})/({lo_hi[2][1]}-{lo_hi[2][0]}),4))",
            "score": f"=ROUND(({A['w1']}*{c('n1')}+{A['w2']}*{c('n2')}+{A['w3']}*{c('n3')})/({A['w1']}+{A['w2']}+{A['w3']}),9)",
            "rank": (
                f'=IF({c("feas")}=1,1+COUNTIFS({rng("feas")},1,{rng("score")},"<"&{c("score")})'
                f'+COUNTIFS(${L["feas"]}${HEAD}:{L["feas"]}{r - 1},1,${L["score"]}${HEAD}:{L["score"]}{r - 1},{c("score")}),"")'
            ),
        }
        return f

    for k, (fshare, pshare, vshare) in enumerate(shares):
        r = FIRST + k
        ws[f"A{r}"] = f"g{fshare}-{pshare}-{vshare}"
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = fshare, pshare, vshare
        fm = formulas(r)
        for i, (key, _, fmt) in enumerate(cols):
            cell = ws.cell(row=r, column=i + 1)
            if fmt:
                cell.number_format = fmt
            if key in fm and (solution or k == 0):
                cell.value = fm[key]
            elif key in fm:
                cell.fill = fill(t["n100"])
                cell.border = box
    ws.freeze_panes = ws[f"B{FIRST}"]

    # results
    rr = lambda key: rng(key)  # noqa: E731
    m = f"MATCH(1,{rr('rank')},0)"
    header(58, "תוצאות החיפוש")
    put(59, "space", "כמה מועמדים ברשת", f"=COUNTA({rr('id')})", plain)
    put(60, "n_feas", "כמה מהם אפשריים", f"=SUM({rr('feas')})", plain)
    put(61, "n_eff", "כמה מהם יעילים", f"=SUM({rr('eff')})", plain)
    put(62, "winner", "המועמד שנבחר", f'=IFERROR(INDEX({rr("id")},{m}),"אין מועמד אפשרי")', None)
    put(63, "w_first", "ההחזר הראשון שלו (₪)", f'=IFERROR(INDEX({rr("first")},{m}),"")', money)
    put(64, "w_total", "סך התשלומים שלו (₪)", f'=IFERROR(INDEX({rr("total")},{m}),"")', money)
    put(65, "w_expo", "החשיפה שלו (₪)", f'=IFERROR(INDEX({rr("expo")},{m}),"")', money)
    m2 = f"MATCH(2,{rr('rank')},0)"
    put(66, "runner", "המקום השני", f'=IFERROR(INDEX({rr("id")},{m2}),"אין")', None)
    put(67, "g_first", "ההפרש בהחזר הראשון: השני פחות הנבחר (₪)", f'=IFERROR(INDEX({rr("first")},{m2})-{A["w_first"]},"")', money)
    put(68, "g_total", "ההפרש בסך התשלומים (₪)", f'=IFERROR(INDEX({rr("total")},{m2})-{A["w_total"]},"")', money)
    put(69, "g_expo", "ההפרש בחשיפה (₪)", f'=IFERROR(INDEX({rr("expo")},{m2})-{A["w_expo"]},"")', money)
    header(71, "למה נפסלו מועמדים: כמה נפסלו בכל אילוץ")
    for k, cid in enumerate(IDS):
        r = 72 + k
        ws[f"A{r}"] = IDS_HE[cid]
        ws[f"B{r}"] = f"=SUM({rng('x_' + cid)})"
        ws[f"A{r}"].font = font(color=t["n800"])
    A["rejects"] = "$B$72:$B$81"
    header(83, "ההסבר, מתוך המספרים בלבד")
    ws["A84"] = (
        f'=IF(ISNUMBER({A["w_first"]}),"נבחר תמהיל "&INDEX({rr("f")},{m})&"% קבועה, "&INDEX({rr("p")},{m})&"% פריים ו־"&INDEX({rr("v")},{m})&"% משתנה. '
        f'ההחזר הראשון "&TEXT({A["w_first"]},"#,##0.00")&" ₪, סך התשלומים "&TEXT({A["w_total"]},"#,##0")&" ₪, החשיפה "&TEXT({A["w_expo"]},"#,##0.00")&" ₪.","אף מועמד לא עבר את כל האילוצים: שנו את הבחירות.")'
    )
    ws["A85"] = (
        f'="נפסלו "&({A["space"]}-{A["n_feas"]})&" מתוך "&{A["space"]}&" מועמדים. מתוך האפשריים "&{A["n_eff"]}&" יעילים: לא ניתן לשפר בהם מדד בלי לפגוע באחר."'
    )
    ws["A86"] = "הניסוח במילים הוא רק ניסוח: כל מספר בו הוא מספר של המנוע. מודל שפה לעולם לא מחשב."
    for r in (84, 85, 86):
        ws[f"A{r}"].font = font(color=t["n800"])
        ws[f"A{r}"].alignment = Alignment(wrap_text=False)

    # an input outside its range cascades through 66 rows, so the green cells refuse it
    from openpyxl.worksheet.datavalidation import DataValidation

    def validate(cell, **kw):
        dv = DataValidation(allow_blank=False, showErrorMessage=True, errorTitle="ערך לא תקין", **kw)
        ws.add_data_validation(dv)
        dv.add(cell)

    validate("B18", type="decimal", operator="between", formula1="0", formula2="500000", error="הון עצמי נוסף: מספר בין 0 ל־500,000")
    validate("B19", type="whole", operator="between", formula1="0", formula2="1", error="0 (לא) או 1 (כן)")
    validate("B20", type="whole", operator="between", formula1="0", formula2="1", error="0 (לא) או 1 (כן)")
    validate("B21", type="list", formula1='"a,b,c,d"', error="בנק: a, b, c או d")
    for cell in ("B22", "B23", "B24"):
        validate(cell, type="whole", operator="between", formula1="0", formula2="100", error="משקולת: מספר שלם בין 0 ל־100")

    # sample offer, checked against the estimate
    offer = case["offer"]
    oc = "F"
    ws["E4"] = "הצעת בנק לדוגמה, מול אומדן"
    ws["E4"].font = font(bold=True, color=navy)
    ob = [
        ("בנק ההצעה", offer["bank"], None),
        ("חלק קבועה %", offer["shares"][0], plain), ("חלק פריים %", offer["shares"][1], plain), ("חלק משתנה %", offer["shares"][2], plain),
        ("ריבית קבועה בהצעה (%)", num(offer["rates"].get("fixed_unlinked", 0)), pct),
        ("ריבית פריים בהצעה (%)", num(offer["rates"].get("prime", 0)), pct),
        ("ריבית משתנה בהצעה (%)", num(offer["rates"].get("variable_unlinked", 0)), pct),
    ]
    for k, (label, v, fmt) in enumerate(ob):
        r = 5 + k
        ws[f"E{r}"] = label
        ws[f"F{r}"] = v
        ws[f"F{r}"].fill = fill(t["mint"])
        ws[f"F{r}"].border = box
        if fmt:
            ws[f"F{r}"].number_format = fmt
    obrow = "MATCH($F$5,$A$28:$A$31,0)"
    n = A["n"]
    base_loan = f"({A['price']}-{A['base_eq']})"
    amt = lambda cell: f"ROUND({base_loan}*{cell}/100,2)"  # noqa: E731
    est_rate = lambda col: f"INDEX(${col}$28:${col}$31,{obrow})"  # noqa: E731
    pay = lambda rate, share: f"IF({share}>0,-PMT({rate}/1200,{n},{amt(share)}),0)"  # noqa: E731
    pay_r = lambda rate, share: f"IF({share}>0,ROUND(-PMT({rate}/1200,{n},{amt(share)}),2),0)"  # noqa: E731
    shares_c = ("$F$6", "$F$7", "$F$8")
    ws["E13"], ws["E14"], ws["E15"], ws["E16"], ws["E17"], ws["E18"] = "אומדן: ההחזר הראשון (₪)", "הצעה: ההחזר הראשון (₪)", "הפער בהחזר הראשון (₪)", "אומדן: סך התשלומים (₪)", "הצעה: סך התשלומים (₪)", "הפער בסך התשלומים (₪)"
    ws["F13"] = "=" + "+".join(pay_r(est_rate(col), s) for col, s in zip("CDE", shares_c))
    ws["F14"] = "=" + "+".join(pay_r(f"$F${9 + j}", s) for j, s in enumerate(shares_c))
    ws["F15"] = "=F14-F13"
    ws["F16"] = "=(" + "+".join(pay(est_rate(col), s) for col, s in zip("CDE", shares_c)) + f")*{n}"
    ws["F17"] = "=(" + "+".join(pay(f"$F${9 + j}", s) for j, s in enumerate(shares_c)) + f")*{n}"
    ws["F18"] = "=F17-F16"
    for r in range(13, 19):
        ws[f"F{r}"].number_format = money
        ws[f"E{r}"].font = font(color=t["n800"])
    ws.column_dimensions["E"].width = 34
    f13 = ws["F13"].value
    if not solution:
        # F13 stays as the worked example; the other five are the learner's
        hints = {
            14: "אותו חישוב כמו ב־F13, אבל עם הריביות של ההצעה (F9 עד F11) במקום המחירון",
            15: "הצעה פחות אומדן: =F14-F13. מינוס אומר שההצעה זולה מהאומדן",
            16: "אותו חישוב כמו ב־F13 בלי ROUND, כפול מספר החודשים",
            17: "אותו חישוב כמו ב־F16, עם הריביות של ההצעה",
            18: "הצעה פחות אומדן: =F17-F16",
        }
        for r in range(14, 19):
            ws[f"F{r}"].value = None
            ws[f"F{r}"].fill = fill(t["n100"])
            ws[f"F{r}"].border = box
            ws[f"F{r}"].comment = Comment(hints[r], "Mogi")
        ws["F13"].comment = Comment("דוגמה: ההחזר הראשון של האומדן. לכל מסלול: ROUND(-PMT(ריבית מהמחירון/1200, חודשים, סכום), 2); אין חלק, אין החזר", "Mogi")

    if not solution:
        notes = {
            "ef": "סכום המסלול = ההלוואה × החלק ÷ 100, מעוגל לאגורה: =ROUND($B$47*B{r}/100,2)",
            "hf": "ההחזר הראשון = ROUND(-PMT(ריבית שנתית/1200, חודשים, סכום), 2); 0 כשאין חלק",
            "qual": "אותו חישוב בלי ROUND, מחובר: ההחזר שיחס ההחזר מודד",
            "total": "ההחזר החודשי הלא־מעוגל כפול מספר החודשים (בסיס: אין תנועה בריבית)",
            "stress": "קבועה: אותו החזר. פריים: מחושב מחדש בחודש 13 עם הריבית המוגדלת על היתרה. משתנה: מחושב מחדש אחרי עדכון הריבית",
            "x_ltv": "1 כשהאילוץ נשבר, 0 כשהוא נשמר. ״מעל״ הוא תמיד חזק",
            "dom": "SUMPRODUCT: כמה מועמדים אפשריים לא גרועים בשום מדד ומשופרים לפחות באחד",
            "score": "ממוצע משוקלל של שלושת המנורמלים: הנמוך מנצח",
            "rank": "1 + כמה אפשריים עם ציון נמוך יותר + כמה לפני בשוויון",
        }
        for key, text in notes.items():
            ws[f"{L[key]}{HEAD}"].comment = Comment(text, "Mogi")
        ws[f"A{HEAD - 1}"] = "שורה 1 של הרשת מלאה כדוגמה. כתבו את אותן נוסחאות בשורות האחרות (העתקה למטה)."
        ws[f"A{HEAD - 1}"].font = font(bold=True, color=t["danger"] if "danger" in t else navy)

    # the check tab
    chk = wb.create_sheet("בדיקה")
    chk.sheet_view.rightToLeft = True
    chk.column_dimensions["A"].width = 52
    for col in "BCDEF":
        chk.column_dimensions[col].width = 20
    chk["A1"] = "הגיליון מול המנוע של הקורס, לקלט ההתחלתי"
    chk["A1"].font = font(size=14, bold=True, color=navy)
    chk["A2"] = "עמודה C היא התוצאה של המנוע של הקורס (מחושבת בקוד, ונבדקת פעמיים). הערכים נכונים רק לקלט ההתחלתי: אם שיניתם בחירה, ההפרשים צפויים."
    chk["A2"].font = font(size=10, color=t["n600"])
    for col, text in zip("ABCDEF", ("מה נבדק", "הגיליון", "המנוע", "הפרש", "סבילות", "תקין?")):
        chk[f"{col}4"] = text
        chk[f"{col}4"].font = font(bold=True, color=white)
        chk[f"{col}4"].fill = fill(navy)
    items = [
        ("כמה מועמדים ברשת", f"='מנוע'!{A['space']}", F["p.space"]["value"], 0),
        ("כמה אפשריים", f"='מנוע'!{A['n_feas']}", F["p.feasible"]["value"], 0),
        ("כמה יעילים", f"='מנוע'!{A['n_eff']}", F["p.efficient"]["value"], 0),
        ("המועמד שנבחר", f"='מנוע'!{A['winner']}", F["p.winner_id"]["value"], None),
        ("ההחזר הראשון של הנבחר", f"='מנוע'!{A['w_first']}", F["p.winner:first"]["value"], 0.05),
        ("סך התשלומים של הנבחר", f"='מנוע'!{A['w_total']}", F["p.winner:total"]["value"], 6),
        ("החשיפה של הנבחר", f"='מנוע'!{A['w_expo']}", F["p.winner:exposure"]["value"], 0.05),
    ]
    for k, cid in enumerate(IDS):
        items.append((f"נפסלו: {IDS_HE[cid]}", f"='מנוע'!$B${72 + k}", F[f"p.reject:{cid}"]["value"], 0))
    items += [
        ("הפער בהחזר הראשון: הצעה מול אומדן", "='מנוע'!$F$15", F["p.offer:first_gap"]["value"], 0.05),
        ("הפער בסך התשלומים: הצעה מול אומדן", "='מנוע'!$F$18", F["p.offer:total_gap"]["value"], 6),
    ]
    for k, (label, formula, expected, tol) in enumerate(items):
        r = 5 + k
        chk[f"A{r}"] = label
        chk[f"B{r}"] = formula
        chk[f"C{r}"] = num(expected) if isinstance(expected, (int, float)) or re.fullmatch(r"-?\d+(\.\d+)?", str(expected)) else expected
        if tol is not None:
            chk[f"D{r}"] = f'=IF(ISNUMBER(B{r}),B{r}-C{r},"")'
            chk[f"E{r}"] = tol
            chk[f"F{r}"] = f'=IF(ISNUMBER(D{r}),IF(ABS(D{r})<=E{r},"✓","✗"),"✗")'
        else:
            chk[f"F{r}"] = f'=IF(B{r}=C{r},"✓","✗")'
        for col in "BCD":
            chk[f"{col}{r}"].number_format = "#,##0.00" if tol not in (0, None) else "0"
    last = 5 + len(items)
    chk[f"A{last + 1}"] = "הגיליון מקרב את המנוע: הוא מחשב החזר מדויק (בלי עיגול כל ריבית חודשית), והמנוע מעגל כל חודש לאגורה. בסך התשלומים הפער עד כמה שקלים, כי במנוע התשלום האחרון סוגר את היתרה והוא שונה בכמה שקלים מהקבוע. בחשיפה אין פער של ממש, כי המנוע לא סופר את התשלום האחרון בהחזר הגבוה. הבחירה של התמהיל לא מושפעת."
    chk[f"A{last + 1}"].font = font(size=10, color=t["n600"])
    chk[f"A{last + 2}"] = f'=IF(COUNTIF(F5:F{last - 1},"✗")=0,"✓ הכול בטווח","✗ יש פערים: בדקו את הנוסחאות או את הקלט")'
    chk[f"A{last + 2}"].font = font(bold=True, color=navy)

    instructions(wb, kind, t, font, L, HEAD, FIRST, LAST, f13)
    wb.properties.title = "מנוע קטן לבחירת תמהיל — סדנה 8.28"
    wb.properties.creator = "Mogi"
    return wb, (L, HEAD, FIRST, LAST, A)


def json_choice(F, key):
    return F["p.loan"]["inputs"][key]


def grid_ids(F):
    # every [fixed, prime, variable] triple in 10-point steps, the fixed share changing slowest (as src/lib/calc/ws8.ts gridShapes)
    step = F["ws.case"]["value"]["rules"]["step"]
    return [f"g{f}-{p}-{100 - f - p}" for f in range(0, 101, step) for p in range(0, 101 - f, step)]


def instructions(wb, kind, t, font, L, HEAD, FIRST, LAST, f13):
    ws = wb.create_sheet("הוראות")
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 120
    cs = lambda k: L[k]  # noqa: E731
    lines = [
        ("h", "איך בונים את המנוע"),
        ("p", MARK),
        ("p", "הגיליון ״מנוע״ מחזיק את התיק, את הבחירות שלכם (תאים ירוקים), את ארבעת הבנקים, את הכללים, ובסוף את הרשת: שורה לכל מועמד, בלי שום מספר שהוקלד ביד בנוסחאות."),
        ("h", "הצעדים"),
        ("p", "1 — מיצוב: בתאים B18 עד B21 בחרו הון עצמי נוסף, סגירת הלוואת הרכב, הורה ובנק. ההלוואה, שיעור המימון והכנסה הפנויה מתעדכנים (שורות 45 עד 56)."),
        ("p", f"2 — תזרים: בכל שורה של הרשת, ההחזר הראשון של כל מסלול (עמודות {cs('hf')} עד {cs('hv')}) הוא ROUND(-PMT(ריבית/1200, חודשים, סכום), 2). סך התשלומים הוא ההחזר הלא־מעוגל כפול החודשים, וההחזר הגבוה בלחץ מחושב מחדש אחרי שהריבית עולה."),
        ("p", f"3 — חיפוש: הרשת היא כל התמהילים בצעדים של 10 אחוזים (66). אין להוסיף שורות: מרחב החיפוש הוא מה שנבדק. שורת הרשת הראשונה ({FIRST}) מלאה בכל העמודות כדוגמה. עמדו על כותרת של עמודה כדי לקרוא מה היא עושה (בכותרות עם הערה), העתיקו את השורה למטה עד שורה {LAST}, ובדקו שהתוצאות בשורות 59 עד 69 התמלאו."),
        ("p", f"4 — אילוצים: עשר עמודות ({cs('x_sum_equals_loan')} עד {cs('x_client_target_ratio')}) של 1 (נשבר) או 0 (נשמר). מועמד אפשרי כשאין בו אף 1. הסיבות מסוכמות בשורות 72 עד 81."),
        ("p", f"5 — חזית יעילה: עמודה {cs('dom')} סופרת כמה מועמדים אפשריים לא גרועים בשום מדד ומשופרים באחד. מועמד אפשרי שאף אחד לא שולט עליו הוא יעיל."),
        ("p", f"6 — בחירה: מנרמלים כל מדד בין אפשריים (0 הטוב, 1 הגרוע), ממוצע משוקלל לפי המשקולות בתאים B22 עד B24 (להדגמה). הציון הנמוך מנצח."),
        ("p", "7 — הסבר: בשורות 84 עד 86, הניסוח נבנה ממספרי הגיליון. מודל שפה לעולם לא מחשב."),
        ("p", "8 — בדיקה מול הצעת בנק: בתאים E4 עד F11 יש הצעה לדוגמה של בנק בדוי. האומדן הוא המחירון של הבנק; ההצעה היא מה שהוא כתב. הפער נמדד בהחזר הראשון ובסך התשלומים."),
        ("p", ("התא F13 מלא כדוגמה, ו־F14 עד F18 ריקים: מלאו אותם בעצמכם. F13 הוא ההחזר הראשון של האומדן: " + f13 + ". ב־F14 אותו חישוב עם ריביות ההצעה (F9 עד F11), F15 הוא F14 פחות F13, ו־F16 עד F18 הם אותו דבר לסך התשלומים: בלי ROUND, כפול מספר החודשים.") if kind == "template" else "התאים F13 עד F18 מלאים: אומדן, הצעה והפער בהחזר הראשון ובסך התשלומים."),
        ("h", "הלשונית ״בדיקה״"),
        ("p", "היא משווה את הגיליון למנוע של הקורס לקלט ההתחלתי. הסבילות היא כמה שקלים, כי הגיליון לא מעגל כל ריבית חודשית. בתבנית היא מתעדכנת בזמן שאתם ממלאים: עמודה C היא התוצאות של המנוע, כך שאפשר להציץ בה בסוף, וכשהכול מתאים יופיע ✓ בתחתית."),
        ("h", "Excel ו־Google Sheets"),
        ("p", "בחלק מההגדרות האזוריות מפרידים בין ארגומנטים בנקודה־פסיק (;). הגיליון משתמש בפונקציות PMT, SUMPRODUCT, COUNTIFS ו־INDEX: הכול קיים גם ב־Google Sheets."),
        ("p", "שינוי בחירה ישנה את התוצאות, וגם את הבדיקה: היא נכונה לקלט ההתחלתי בלבד."),
    ]
    for i, (style, text) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = font(bold=True, size=13, color=t["navy"]) if style == "h" else font(color=t["n800"])


# ── reproducible output with stored values ─────────────────────────────────


def recalc(xlsx, workdir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("needs LibreOffice (soffice) on the PATH")
    out = workdir / "out"
    out.mkdir(exist_ok=True)
    profile = (workdir / "profile").as_uri()
    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(xlsx)], check=True, capture_output=True, timeout=240)
    return out / xlsx.name


def with_cached(xml, cached):
    def fill(m):
        ref, attrs, formula = m.group(1), m.group(2), m.group(3)
        value = cached.get(ref)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return f'<c r="{ref}"{attrs}><f>{formula}</f><v>{repr(float(value)) if isinstance(value, float) else value}</v></c>'
        text = "" if value is None else str(value)
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return f'<c r="{ref}"{attrs} t="str"><f>{formula}</f><v>{text}</v></c>'

    return re.sub(r'<c r="([A-Z]+[0-9]+)"((?: s="[0-9]+")?)><f>(.*?)</f><v></v></c>', fill, xml)


def cached_values(xlsx):
    with tempfile.TemporaryDirectory() as d:
        out = recalc(xlsx, Path(d))
        wf = load_workbook(xlsx)
        wv = load_workbook(out, data_only=True)
        return [{c.coordinate: wv[name][c.coordinate].value for row in wf[name].iter_rows() for c in row if c.data_type == "f"} for name in wf.sheetnames]


def save(wb, path, cached=None):
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
            if cached and m and int(m.group(1)) - 1 < len(cached):
                data = with_cached(data.decode("utf-8"), cached[int(m.group(1)) - 1]).encode("utf-8")
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", rf"\g<1>{FIXED_DATE}\g<2>", text)
                data = text.encode("utf-8")
            out = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            out.compress_type = zipfile.ZIP_DEFLATED
            out.external_attr = 0o644 << 16
            dst.writestr(out, data)
    tmp.unlink()


# ── the notebook ───────────────────────────────────────────────────────────


def notebook(F):
    """checks/ws8_engine.py is in jupytext's percent format: '# %% [markdown]' cells (comment lines) and '# %%' code cells."""
    src = (ROOT / "checks" / "ws8_engine.py").read_text(encoding="utf-8").splitlines()
    cells, cur, kind = [], [], None

    def flush():
        if kind is None:
            return
        if kind == "markdown":
            text = "\n".join(re.sub(r"^# ?", "", ln) for ln in cur).strip()
            cells.append({"cell_type": "markdown", "metadata": {}, "source": text.splitlines(keepends=True)})
        else:
            text = "\n".join(cur).strip("\n")
            cells.append({"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [], "source": text.splitlines(keepends=True)})

    for line in src:
        if line.startswith("# %% [markdown]"):
            flush()
            cur, kind = [], "markdown"
        elif line.startswith("# %%"):
            flush()
            cur, kind = [], "code"
        else:
            cur.append(line)
    flush()

    case = F["ws.case"]["value"]
    inp = F["p.loan"]["inputs"]
    choice = {"extra_equity": str(inp["extra_equity"]), "close_car": bool(int(inp["close_car"])), "parent": bool(int(inp["parent"])), "bank": inp["bank"]}
    weights = [int(x) for x in inp["weights"].split("|")]
    case_src = json.dumps(case, ensure_ascii=False, indent=1)
    inputs_md = {
        "cell_type": "markdown",
        "metadata": {},
        "source": (
            "## נתוני התיק והבחירות שלכם\n\n"
            "* `CASE` הוא התיק הסינתטי, כפי שהוא באתר: מחיר, הון עצמי, הכנסות, הלוואת רכב, ארבעה בנקים בדויים וכללים. אין בו מספר שהוקלד ביד: הוא נוצר מאותם נתונים שמהם נוצר הגיליון.\n"
            "* `CHOICE` הוא המיצוב שלכם: הון עצמי נוסף, סגירת הלוואת הרכב, הורה כלווה נוסף, ובנק (`a` עד `d`).\n"
            "* `WEIGHTS` הם המשקולות של שלושת המדדים (החזר ראשון, סך התשלומים, חשיפה). הם **להדגמה**.\n\n"
            "שנו את `CHOICE` או את `WEIGHTS`, והריצו מחדש את התאים מהתא של ״להריץ הכול״.\n"
        ).splitlines(keepends=True),
    }
    inputs_code = {
        "cell_type": "code",
        "metadata": {},
        "execution_count": None,
        "outputs": [],
        "source": (f"CASE = {case_src}\n\nCHOICE = {choice!r}\nWEIGHTS = {weights!r}\n").splitlines(keepends=True),
    }
    tail_md = {
        "cell_type": "markdown",
        "metadata": {},
        "source": (
            "## להריץ ולקרוא את התוצאה\n\n"
            "התא הבא מריץ את כל השרשרת, מדפיס את המועמדים האפשריים ואת מי שיעיל, בוחר, ומסביר ממספרי המנוע. "
            "אחרי זה הוא משווה את האומדן להצעת הבנק הלדוגמה. לבסוף, `EXPECTED` הם המספרים שהאתר מציג לקלט ההתחלתי, "
            "והמחברת בודקת שהקוד שלכם מגיע לאותם מספרים אם לא שיניתם את הקלט.\n"
        ).splitlines(keepends=True),
    }
    expected = {
        "feasible": F["p.feasible"]["value"],
        "efficient": F["p.efficient"]["value"],
        "winner": F["p.winner_id"]["value"],
        "winner_first": F["p.winner:first"]["value"],
        "winner_total": F["p.winner:total"]["value"],
        "winner_exposure": F["p.winner:exposure"]["value"],
        "offer_first_gap": F["p.offer:first_gap"]["value"],
        "offer_total_gap": F["p.offer:total_gap"]["value"],
    }
    run_code = {
        "cell_type": "code",
        "metadata": {},
        "execution_count": None,
        "outputs": [],
        "source": (
            "result = run(CASE, CHOICE, WEIGHTS)\n"
            "pos = result['positioning']\n"
            "print('הלוואה:', agorot(pos['loan']), '  הכנסה פנויה:', agorot(pos['disposable']))\n"
            "rows = result['rows']\n"
            "print('מועמדים:', len(rows), ' אפשריים:', sum(r['feasible'] for r in rows), ' יעילים:', sum(r['efficient'] for r in rows))\n"
            "print()\n"
            "print('המועמדים האפשריים (קבועה, פריים, משתנה | החזר ראשון, סך, חשיפה | יעיל? | מקום):')\n"
            "for r in rows:\n"
            "    if r['feasible']:\n"
            "        print(r['shares'], [str(agorot(v)) for v in r['values']], r['efficient'], r['rank'])\n"
            "print()\n"
            "print('מועמד שנפסל, עם כל הסיבות שלו (הראשון בטבלה שנפסל ביותר מסיבה אחת):')\n"
            "for r in rows:\n"
            "    if len(r['reasons']) > 1:\n"
            "        print(r['id'], r['reasons'])\n"
            "        break\n"
            "print()\n"
            "ex = explain(result)\n"
            "if ex:\n"
            "    w = ex['winner']\n"
            "    print('נבחר: %d%% קבועה, %d%% פריים, %d%% משתנה.' % w['shares'])\n"
            "    print('ההחזר הראשון', agorot(w['values'][0]), '₪, סך התשלומים', agorot(w['values'][1]), '₪, חשיפה', agorot(w['values'][2]), '₪.')\n"
            "    if ex['gaps']:\n"
            "        print('המקום השני יקר ממנו ב', [str(agorot(g)) for g in ex['gaps']], '(החזר ראשון, סך, חשיפה).')\n"
            "else:\n"
            "    print('אף מועמד לא עבר את כל האילוצים: שנו את CHOICE.')\n"
            "print()\n"
            "offer = check_offer(CASE)\n"
            "print('אומדן מול הצעת בנק: הפער בהחזר הראשון', agorot(offer['first_gap']), '₪, בסך התשלומים', agorot(offer['total_gap']), '₪.')\n"
        ).splitlines(keepends=True),
    }
    check_code = {
        "cell_type": "code",
        "metadata": {},
        "execution_count": None,
        "outputs": [],
        "source": (
            f"EXPECTED = {json.dumps(expected, ensure_ascii=False, indent=1)}\n\n"
            "if CHOICE == {'extra_equity': '%s', 'close_car': %s, 'parent': %s, 'bank': '%s'} and WEIGHTS == %s:\n"
            "    assert sum(r['feasible'] for r in rows) == EXPECTED['feasible']\n"
            "    assert sum(r['efficient'] for r in rows) == EXPECTED['efficient']\n"
            "    assert ex['winner']['id'] == EXPECTED['winner']\n"
            "    assert str(agorot(ex['winner']['values'][0])) == EXPECTED['winner_first']\n"
            "    assert str(agorot(ex['winner']['values'][1])) == EXPECTED['winner_total']\n"
            "    assert str(agorot(ex['winner']['values'][2])) == EXPECTED['winner_exposure']\n"
            "    assert str(agorot(offer['first_gap'])) == EXPECTED['offer_first_gap']\n"
            "    assert str(agorot(offer['total_gap'])) == EXPECTED['offer_total_gap']\n"
            "    print('✓ המחברת מגיעה לאותם מספרים כמו האתר')\n"
            "else:\n"
            "    print('שיניתם את הקלט, ולכן המספרים שונים מהאתר. זה צפוי.')\n"
        % (choice["extra_equity"], choice["close_car"], choice["parent"], choice["bank"], json.dumps(weights))
        ).splitlines(keepends=True),
    }
    # the first markdown cell (intro), then the inputs, then the engine's cells, then the run
    out = cells[:1] + [inputs_md, inputs_code] + cells[1:] + [tail_md, run_code, check_code]
    nb = {
        "cells": out,
        "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}, "language_info": {"name": "python"}},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    for i, c in enumerate(nb["cells"]):
        c["id"] = f"c{i:02d}"
    return nb


# ── --verify ───────────────────────────────────────────────────────────────

CHOICES = [
    ("default", {}),
    ("no extra equity", {"B18": 0}),
    ("parent joins, bank d", {"B20": 1, "B21": "d", "B18": 0}),
    ("bank c, heavy weight on exposure", {"B21": "c", "B22": 10, "B23": 20, "B24": 70}),
    ("car closed with extra equity (nothing feasible)", {"B19": 1}),
]


def verify(solution, F):
    import ws8_engine as eng

    case = F["ws.case"]["value"]
    problems = []
    inp = F["p.loan"]["inputs"]
    base_choice = {"extra_equity": str(inp["extra_equity"]), "close_car": bool(int(inp["close_car"])), "parent": bool(int(inp["parent"])), "bank": inp["bank"]}
    base_weights = [int(x) for x in inp["weights"].split("|")]
    worst = {"first": Decimal(0), "total": Decimal(0), "expo": Decimal(0)}
    with tempfile.TemporaryDirectory() as tmp:
        for name, change in CHOICES:
            wb = load_workbook(solution)
            ws = wb["מנוע"]
            for cell, v in change.items():
                ws[cell] = v
            d = Path(tmp) / re.sub(r"\W+", "_", name)
            d.mkdir()
            src = d / "case.xlsx"
            wb.save(src)
            got = load_workbook(recalc(src, d), data_only=True)["מנוע"]
            choice = {"extra_equity": str(int(got["B18"].value)), "close_car": bool(got["B19"].value), "parent": bool(got["B20"].value), "bank": got["B21"].value}
            weights = [int(got["B22"].value), int(got["B23"].value), int(got["B24"].value)]
            result = eng.run(case, choice, weights)
            rows = result["rows"]
            HEAD = 88
            # find the columns by the header text
            heads = {ws.cell(row=HEAD, column=i).value: i for i in range(1, 45)}
            col = lambda label: heads[label]  # noqa: E731
            bad = 0
            swaps = 0
            for k, r in enumerate(rows):
                gr = HEAD + 1 + k
                feas = got.cell(row=gr, column=col("אפשרי? (1 כן)")).value
                eff = got.cell(row=gr, column=col("יעיל? (1 כן)")).value
                rank = got.cell(row=gr, column=col("דירוג בין האפשריים")).value
                if bool(feas) != r["feasible"]:
                    bad += 1
                    problems.append(f"{name}: {r['id']} feasible: sheet {feas}, engine {r['feasible']}, reasons {r['reasons']}")
                if bool(eff) != bool(r["efficient"]):
                    bad += 1
                    problems.append(f"{name}: {r['id']} efficient: sheet {eff}, engine {r['efficient']}")
                if (rank or None) != r.get("rank"):
                    # The sheet is an approximation: rows whose scores are within rounding of each other may swap places.
                    # The winner and the runner-up (the two the page names) must agree.
                    swaps += 1
                    if (rank or 99) <= 2 or (r.get("rank") or 99) <= 2:
                        bad += 1
                        problems.append(f"{name}: {r['id']} rank: sheet {rank}, engine {r.get('rank')}")
                for key, label, idx in (("first", "ההחזר הראשון", 0), ("total", "סך התשלומים", 1), ("expo", "חשיפה", 2)):
                    sv = got.cell(row=gr, column=col(label)).value
                    diff = abs(Decimal(repr(sv)) - r["values"][idx])
                    worst[key] = max(worst[key], diff)
            winner = next((r for r in rows if r.get("rank") == 1), None)
            print(f"  {name}: {len(rows)} rows compared; feasible {sum(r['feasible'] for r in rows)}, winner {winner['id'] if winner else None}; sheet says {got['B62'].value}; places swapped by rounding below the runner-up: {swaps}")
            if (got["B62"].value if winner else "אין מועמד אפשרי") != (winner["id"] if winner else "אין מועמד אפשרי"):
                problems.append(f"{name}: winner differs: sheet {got['B62'].value}, engine {winner and winner['id']}")
    print(f"  largest gap, sheet against the engine: first payment {worst['first']}, total {worst['total']}, exposure {worst['expo']}")
    if worst["first"] > Decimal("0.05") or worst["total"] > 6 or worst["expo"] > Decimal("0.05"):
        problems.append(f"gap larger than the tolerance of the check tab: {worst}")
    return problems


def verify_notebook(nb, F):
    """Run the notebook's code cells in a fresh namespace, as a learner would from the top; the EXPECTED asserts must pass."""
    ns = {"__name__": "notebook"}
    import contextlib
    import io

    for c in nb["cells"]:
        if c["cell_type"] == "code":
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                exec("".join(c["source"]), ns)  # noqa: S102
            if "✓" in buf.getvalue():
                print("  notebook:", buf.getvalue().strip().splitlines()[-1])
                return []
    return ["the notebook's last cell did not report agreement with the page"]


def main():
    if not DATA.exists():
        raise SystemExit("data/lessons/8-28.json is missing; run pnpm data first")
    F = facts()
    OUT.mkdir(parents=True, exist_ok=True)
    have_office = bool(shutil.which("soffice") or shutil.which("libreoffice"))
    paths = {}
    for kind in ("template", "solution"):
        paths[kind] = OUT / f"engine-workshop-{kind}.xlsx"
        save(build(kind, F)[0], paths[kind])
        if have_office:
            save(build(kind, F)[0], paths[kind], cached_values(paths[kind]))
        print(f"wrote {paths[kind].relative_to(ROOT)}" + ("" if have_office else " (no LibreOffice: formulas without stored values)"))
    nb = notebook(F)
    nbp = OUT / "engine-workshop.ipynb"
    nbp.write_text(json.dumps(nb, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {nbp.relative_to(ROOT)}")
    if "--verify" in sys.argv:
        print("verifying in LibreOffice and in Python:")
        problems = verify(paths["solution"], F) + verify_notebook(nb, F)
        if problems:
            print("\n".join(f"  ✗ {p}" for p in problems[:40]))
            return 1
        print("  ✓ the solution sheet agrees with the engine on feasibility, front and winner; the notebook reaches the page's numbers")
    return 0


if __name__ == "__main__":
    sys.exit(main())
