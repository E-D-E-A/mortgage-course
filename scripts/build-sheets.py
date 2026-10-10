"""Builds the spreadsheets of workshop 3.12 (a level-payment schedule) into public/downloads/.

    python3 scripts/build-sheets.py            # writes the two files
    python3 scripts/build-sheets.py --verify   # also recalculates them in LibreOffice and checks the numbers

Two files, the same layout (the page names these cells; change both together):
  schedule-template.xlsx  inputs, headers, month numbers 1–360 and the check cells; the
                          learner writes the formulas of the monthly rate, the payment and
                          rows 20–21, then copies down.
  schedule-solution.xlsx  every formula filled, each row guarded so it stays empty after
                          the last month (copy-down to 360 months works for any term).

    A1 title · A2 marking
    B5 loan (₪) · B6 annual rate · B7 months                      (inputs)
    B10 monthly rate =B6/12 · B11 payment =ROUND(-PMT(B10,B7,B5),2)
    B13 last payment · B14 total paid · B15 total interest · B16 total principal · B17 check
    row 19 headers; rows 20–379 months 1–360:
      A month · B opening · C payment · D interest · E principal · F closing

The convention is the course's statement convention (src/lib/calc/amortization.ts,
annuityStatement): the payment rounded once, each interest rounded, the last payment closes.
The default inputs, and the figures --verify expects, are read from data/lessons/3-12.json
(run `pnpm data` first), so the sheet and the page share one source. Colours are read from
src/styles/tokens.css. The output is byte-for-byte reproducible: fixed document dates and
fixed zip timestamps.

--verify needs LibreOffice (soffice) on the PATH. It recalculates the solution three times — with
its default inputs (the hand loan), with the sheet loan of the page, and with a loan whose first
interest is exactly half an agora (a rounding tie) — reads the computed
cells back with openpyxl, and compares every row and total with an independent Python
Decimal recomputation and with the facts the page shows. It also checks that the template
is empty where the learner writes and that its check cell says so.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public" / "downloads"
DATA = ROOT / "data" / "lessons" / "3-12.json"
TOKENS = ROOT / "src" / "styles" / "tokens.css"

FIRST, LAST_MONTH = 20, 360
LAST = FIRST + LAST_MONTH - 1  # row 379
RANGE = lambda col: f"{col}{FIRST}:{col}{LAST}"  # noqa: E731
FIXED_DATE = "2026-10-10T00:00:00Z"
ZIP_TIME = (2026, 10, 10, 0, 0, 0)

MARK = "תרגיל נפרד, לא התיק של נועם ושירה. הלוואות בדויות; המספרים להמחשה, לא נתוני שוק ולא הצעה של בנק."


def facts():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    return {f["id"]: f for f in data["facts"]}


def tokens():
    css = TOKENS.read_text(encoding="utf-8")
    return {m.group(1): m.group(2).upper() for m in re.finditer(r"--mogi-([a-z0-9-]+):\s*#([0-9A-Fa-f]{6})", css)}


# ── the layout ─────────────────────────────────────────────────────────────


def build(kind, loan, rate, months):
    """kind: "template" or "solution". loan, rate (a fraction, 0.06), months: the default inputs."""
    t = tokens()
    navy, white = t["navy"], t["n0"]
    font = lambda **kw: Font(name="Arial", size=kw.pop("size", 11), **kw)  # noqa: E731
    fill = lambda hex_: PatternFill("solid", fgColor=hex_)  # noqa: E731
    thin = Side(style="thin", color=t["n200"])
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    money, pct, pct4 = "#,##0.00", "0.00%", "0.0000%"

    wb = Workbook()
    ws = wb.active
    ws.title = "לוח סילוקין"
    ws.sheet_view.rightToLeft = True
    for col, width in zip("ABCDEF", (30, 22, 16, 14, 16, 20)):
        ws.column_dimensions[col].width = width

    ws["A1"] = "לוח סילוקין בשפיצר — סדנה 3.12" + (" (פתור)" if kind == "solution" else "")
    ws["A1"].font = font(size=14, bold=True, color=navy)
    ws["A2"] = MARK
    ws["A2"].font = font(size=10, color=t["n600"])

    def header(cell, text):
        ws[cell] = text
        ws[cell].font = font(bold=True, color=navy)

    def label(row, text, note=None):
        ws[f"A{row}"] = text
        ws[f"A{row}"].font = font(color=t["n800"])
        if note:
            ws[f"C{row}"] = note
            ws[f"C{row}"].font = font(size=10, color=t["n600"])

    header("A4", "קלט — שנו כאן")
    label(5, "סכום ההלוואה (₪)")
    label(6, "אחוז הריבית לשנה (בחוזה)", "עם סימן אחוז: 6%")
    label(7, "מספר החודשים", f"עד {LAST_MONTH}")
    ws["B5"], ws["B6"], ws["B7"] = loan, rate, months
    for c, fmt in (("B5", money), ("B6", pct), ("B7", "0")):
        ws[c].number_format = fmt
        ws[c].fill = fill(t["mint"])
        ws[c].border = box
        ws[c].font = font(bold=True, color=navy)

    header("A9", "חישוב")
    label(10, "הריבית החודשית", "הריבית לשנה חלקי 12" if kind == "solution" else "את הנוסחה כותבים בתא B10 (צעד 1), לא בתא הזה")
    label(11, "ההחזר החודשי, מעוגל לאגורה", "מעוגל לאגורה" if kind == "solution" else "את הנוסחה כותבים בתא B11 (צעד 2), לא בתא הזה")
    ws["B10"].number_format, ws["B11"].number_format = pct4, money
    if kind == "solution":
        ws["B10"] = "=B6/12"
        ws["B11"] = "=ROUND(-PMT(B10,B7,B5),2)"
    else:
        for c in ("B10", "B11"):
            ws[c].fill = fill(t["n100"])
            ws[c].border = box

    header("A12", "סיכום ובדיקה — מתמלא לבד")
    label(13, "ההחזר בחודש האחרון")
    label(14, "סך התשלומים")
    label(15, "סך הריבית")
    label(16, "סך הקרן שהוחזרה")
    label(17, "בדיקה")
    ws["B13"] = f'=IF(COUNT({RANGE("C")})=0,"",INDEX({RANGE("C")},B7))'
    ws["B14"] = f"=SUM({RANGE('C')})"
    ws["B15"] = f"=SUM({RANGE('D')})"
    ws["B16"] = f"=SUM({RANGE('E')})"
    ws["B17"] = (
        f'=IF(COUNT({RANGE("F")})=0,"עוד לא מילאתם את הלוח",'
        f'IF(AND(ABS(B16-B5)<0.005,ABS(INDEX({RANGE("F")},B7))<0.005),'
        '"✓ הלוח סוגר: הקרן שהוחזרה שווה להלוואה, והיתרה בסוף היא 0",'
        '"✗ הלוח לא סוגר. בדקו את הנוסחאות, ושהעתקתם בדיוק עד החודש האחרון"))'
    )
    for c in ("B13", "B14", "B15", "B16"):
        ws[c].number_format = money
    ws["B17"].fill = fill(t["n50"])
    ws["B17"].font = font(bold=True, color=navy)
    ws["B17"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B17:F17")
    ws.row_dimensions[17].height = 30

    heads = ("חודש", "יתרה בתחילת החודש", "תשלום", "ריבית", "קרן", "יתרה בסוף החודש")
    for col, text in zip("ABCDEF", heads):
        c = ws[f"{col}19"]
        c.value = text
        c.font = font(bold=True, color=white)
        c.fill = fill(navy)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[19].height = 30
    ws.freeze_panes = "A20"
    # Printing: all six columns on one page width, the header row on every page.
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "19:19"

    for k in range(1, LAST_MONTH + 1):
        r = FIRST + k - 1
        ws[f"A{r}"] = k
        ws[f"A{r}"].font = font(color=t["n700"])
        ws[f"A{r}"].alignment = Alignment(horizontal="center")
        for col in "BCDEF":
            ws[f"{col}{r}"].number_format = money
        if kind == "solution":
            guard = lambda body: f'=IF(A{r}>$B$7,"",{body})'  # noqa: E731
            ws[f"B{r}"] = "=B5" if k == 1 else guard(f"F{r - 1}")
            ws[f"C{r}"] = guard(f"IF(A{r}=$B$7,B{r}+D{r},$B$11)")
            ws[f"D{r}"] = guard(f"ROUND(B{r}*$B$10,2)")
            ws[f"E{r}"] = guard(f"C{r}-D{r}")
            ws[f"F{r}"] = guard(f"B{r}-E{r}")
        elif k <= 2:
            for col in "BCDEF":
                ws[f"{col}{r}"].fill = fill(t["n100"])
                ws[f"{col}{r}"].border = box

    if kind == "template":
        ws["B20"].comment = Comment("החודש הראשון: היתרה היא סכום ההלוואה (‎=B5‎).", "Mogi")
        ws["B21"].comment = Comment("מהחודש השני: היתרה בסוף החודש הקודם (‎=F20‎). אחר כך מעתיקים את שורה 21 למטה.", "Mogi")

    instructions(wb, kind, t, font)
    wb.properties.title = "לוח סילוקין בשפיצר — סדנה 3.12"
    wb.properties.creator = "Mogi course (scripts/build-sheets.py)"
    return wb


def instructions(wb, kind, t, font):
    ws = wb.create_sheet("הוראות")
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 110
    lines = [
        ("h", "איך ממלאים את הלוח"),
        ("p", MARK),
        ("p", "הכלל של הקורס: התשלום מעוגל לאגורה פעם אחת; הריבית של כל חודש מעוגלת לאגורה; בחודש האחרון התשלום הוא היתרה ועוד הריבית, כדי שהיתרה תגיע לאפס בדיוק."),
        ("h", "הנוסחאות (גיליון ״לוח סילוקין״)"),
        ("p", "צעד 1 — הריבית החודשית, בתא B10:   =B6/12"),
        ("p", "צעד 2 — ההחזר, בתא B11:   =ROUND(-PMT(B10,B7,B5),2)   — PMT מחזירה מספר שלילי (תשלום יוצא), ולכן המינוס לפניה."),
        ("p", "צעד 3 — החודש הראשון, שורה 20:   B20 =B5   ·   C20 =IF(A20=$B$7,B20+D20,$B$11)   ·   D20 =ROUND(B20*$B$10,2)   ·   E20 =C20-D20   ·   F20 =B20-E20"),
        ("p", "צעד 4 — החודש השני, שורה 21:   B21 =F20, ואת C20 עד F20 מעתיקים ל־C21 עד F21."),
        ("p", "אחר כך מסמנים את B21 עד F21 ומעתיקים למטה, בדיוק עד השורה של החודש האחרון (המספר בעמודה A)."),
        ("p", "צעד 5 — הבדיקה בתא B17 מראה ✓ כשסך הקרן שווה להלוואה והיתרה בסוף החודש האחרון היא 0."),
        ("h", "Excel ו־Google Sheets"),
        ("p", "בחלק מההגדרות האזוריות מפרידים בין הארגומנטים בנקודה־פסיק:   =ROUND(B20*$B$10;2)"),
        ("p", "בתא הריבית כתבו 6% עם סימן האחוז, או 0.06. בלי הסימן, 6 הוא 600%."),
        ("p", "ב־Google Sheets: מעלים את הקובץ ל־Google Drive ופותחים אותו; הנוסחאות נשארות."),
        ("h", "ההשוואה"),
        ("p", "בסדנה 3.12 באתר הקורס מכניסים את ההחזר בחודש 1, ההחזר בחודש האחרון, סך התשלומים וסך הריבית (התאים B11, B13, B14, B15), והדף משווה אותם למחשבון של הקורס."),
    ]
    if kind == "solution":
        lines.insert(
            9,
            ("p", "בגיליון הפתור כל נוסחה בלוח עטופה ב־IF נוסף: אם מספר החודש (עמודה A) גדול ממספר החודשים (B7), השורה נשארת ריקה. כך הלוח עובד לכל תקופה עד 360 חודשים בלי להעתיק מחדש."),
        )
    for i, (style, text) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.data_type = "s"  # a line may start with "=" in another sheet; here it is always text
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = font(bold=True, size=13, color=t["navy"]) if style == "h" else font(color=t["n800"])


# ── reproducible output ────────────────────────────────────────────────────


def save(wb, path):
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", rf"\g<1>{FIXED_DATE}\g<2>", text)
                data = text.encode("utf-8")
            out = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            out.compress_type = zipfile.ZIP_DEFLATED
            out.external_attr = 0o644 << 16
            dst.writestr(out, data)
    tmp.unlink()


# ── independent recomputation, for --verify ────────────────────────────────

AG = Decimal("0.01")


def ag(x):
    return Decimal(x).quantize(AG, rounding=ROUND_HALF_UP)


def statement(principal, annual_percent, months):
    """Rows (opening, payment, interest, principal, closing) by the course's convention, in plain Decimal."""
    m = Decimal(annual_percent) / 12 / 100
    p = Decimal(principal)
    factor = Decimal(1)
    for _ in range(months):
        factor /= 1 + m
    pay = ag(p * m / (1 - factor)) if m else ag(p / months)
    b, rows = ag(p), []
    for k in range(1, months + 1):
        i = ag(b * m)
        part = b if k == months else pay - i
        rows.append((b, part + i, i, part, b - part))
        b -= part
    return rows


def recalc(xlsx, workdir):
    """Recalculate in LibreOffice and return the recalculated copy's path."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("--verify needs LibreOffice (soffice) on the PATH")
    out = workdir / "out"
    out.mkdir(exist_ok=True)
    profile = (workdir / "profile").as_uri()
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(xlsx)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    return out / xlsx.name


def read_back(path, months):
    ws = load_workbook(path, data_only=True)["לוח סילוקין"]
    rows = []
    for k in range(1, months + 1):
        r = FIRST + k - 1
        rows.append(tuple(ws[f"{c}{r}"].value for c in "BCDEF"))
    after = [ws[f"{c}{FIRST + months}"].value for c in "BCDEF"] if months < LAST_MONTH else []
    summary = {c: ws[c].value for c in ("B10", "B11", "B13", "B14", "B15", "B16", "B17")}
    return rows, after, summary


def same(sheet_value, expected):
    return sheet_value is not None and not isinstance(sheet_value, str) and ag(Decimal(repr(sheet_value))) == ag(expected)


def fill_as_taught(ws, months):
    """Write into the template what the page tells the learner to write: B10, B11, rows 20 and 21,
    then row 21 copied down to the last month (a copy shifts the unlocked row numbers)."""
    ws["B10"] = "=B6/12"
    ws["B11"] = "=ROUND(-PMT(B10,B7,B5),2)"
    for k in range(1, months + 1):
        r = FIRST + k - 1
        ws[f"B{r}"] = "=B5" if k == 1 else f"=F{r - 1}"
        ws[f"C{r}"] = f"=IF(A{r}=$B$7,B{r}+D{r},$B$11)"
        ws[f"D{r}"] = f"=ROUND(B{r}*$B$10,2)"
        ws[f"E{r}"] = f"=C{r}-D{r}"
        ws[f"F{r}"] = f"=B{r}-E{r}"


def verify_case(name, xlsx, workdir, loan, rate_percent, months, expect_facts, as_taught=False):
    """expect_facts: {summary cell: expected Decimal string} from data/lessons/3-12.json.
    as_taught: xlsx is the template, filled with the page's formulas before recalculating."""
    problems = []
    case = workdir / name
    case.mkdir()
    wb = load_workbook(xlsx)
    ws = wb["לוח סילוקין"]
    ws["B5"], ws["B6"], ws["B7"] = float(loan), float(Decimal(rate_percent) / 100), months
    if as_taught:
        fill_as_taught(ws, months)
    src = case / "solution.xlsx"
    wb.save(src)
    rows, after, summary = read_back(recalc(src, case), months)
    expected = statement(loan, rate_percent, months)
    for k, (got, want) in enumerate(zip(rows, expected), start=1):
        for col, g, w in zip("BCDEF", got, want):
            if not same(g, w):
                problems.append(f"{name}: month {k} column {col}: sheet {g!r}, expected {w}")
    if any(v not in (None, "") for v in after):
        problems.append(f"{name}: the row after the last month is not empty: {after}")
    for cell, want in expect_facts.items():
        if not same(summary[cell], Decimal(want)):
            problems.append(f"{name}: {cell} is {summary[cell]!r}, the page's fact says {want}")
    if not str(summary["B17"]).startswith("✓"):
        problems.append(f"{name}: the check cell says {summary['B17']!r}")
    print(f"  {name}: {months} rows × 5 columns and {len(expect_facts)} totals compared; check cell: {summary['B17']}")
    return problems


def verify(template, solution, F):
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp)
        # The hand loan: the solution's own defaults. Every row against the page's schedule fact, too.
        h = F["hand.schedule"]["inputs"]
        hand_rows = F["hand.schedule"]["value"]
        problems += verify_case(
            "hand loan",
            solution,
            workdir,
            h["principal"],
            h["rate_percent"],
            int(h["months"]),
            {"B11": F["hand.payment"]["value"], "B13": F["hand.last_payment"]["value"], "B14": F["hand.total_paid"]["value"], "B15": F["hand.total_interest"]["value"], "B16": F["hand.loan"]["value"]},
        )
        mine = statement(h["principal"], h["rate_percent"], int(h["months"]))
        for row, fact in zip(mine, hand_rows):
            if [str(x) for x in row] != [fact["opening"], fact["payment"], fact["interest"], fact["principal_paid"], fact["closing"]]:
                problems.append(f"hand loan: the Python recomputation and the page's fact differ in month {fact['period']}")
        # The sheet loan of the page's step 3 (lesson 2.9's worked example).
        s = F["sheet.statement.total_paid"]["inputs"]
        problems += verify_case(
            "sheet loan",
            solution,
            workdir,
            s["principal"],
            s["rate_percent"],
            int(s["months"]),
            {
                "B11": F["sheet.statement.regular"]["value"],
                "B13": F["sheet.statement.last"]["value"],
                "B14": F["sheet.statement.total_paid"]["value"],
                "B15": F["sheet.statement.total_interest"]["value"],
                "B16": F["sheet.loan"]["value"],
            },
        )
        # A loan whose first interest is exactly half an agora (10,010 × 0.45% = 45.045): the sheet's
        # ROUND must round it up, as the course does, though 45.045 has no exact binary form.
        problems += verify_case("half-agora loan", solution, workdir, "10010", "5.4", 24, {"B16": "10010"})
        # The template, filled with exactly the formulas the page teaches, for both loans.
        problems += verify_case("template as taught, hand loan", template, workdir, h["principal"], h["rate_percent"], int(h["months"]), {"B14": F["hand.total_paid"]["value"], "B15": F["hand.total_interest"]["value"]}, as_taught=True)
        problems += verify_case("template as taught, sheet loan", template, workdir, s["principal"], s["rate_percent"], int(s["months"]), {"B14": F["sheet.statement.total_paid"]["value"], "B15": F["sheet.statement.total_interest"]["value"]}, as_taught=True)
        # The template: empty where the learner writes; its check cell says so.
        wb = load_workbook(template)
        ws = wb["לוח סילוקין"]
        filled = [c for c in ("B10", "B11") + tuple(f"{col}{r}" for r in range(FIRST, LAST + 1) for col in "BCDEF") if ws[c].value is not None]
        if filled:
            problems.append(f"template: cells the learner writes are already filled: {filled[:5]}")
        tcase = workdir / "template"
        tcase.mkdir()
        _, _, summary = read_back(recalc(template, tcase), 1)
        if summary["B17"] != "עוד לא מילאתם את הלוח":
            problems.append(f"template: the check cell says {summary['B17']!r}")
        print(f"  template: formula cells empty; check cell: {summary['B17']}")
    return problems


def main():
    if not DATA.exists():
        raise SystemExit("data/lessons/3-12.json is missing; run pnpm data first")
    F = facts()
    h = F["hand.schedule"]["inputs"]
    loan, rate, months = int(Decimal(h["principal"])), float(Decimal(h["rate_percent"]) / 100), int(h["months"])
    OUT.mkdir(parents=True, exist_ok=True)
    paths = {}
    for kind in ("template", "solution"):
        paths[kind] = OUT / f"schedule-{kind}.xlsx"
        save(build(kind, loan, rate, months), paths[kind])
        print(f"wrote {paths[kind].relative_to(ROOT)}")
    if "--verify" in sys.argv:
        print("verifying in LibreOffice:")
        problems = verify(paths["template"], paths["solution"], F)
        if problems:
            print("\n".join(f"  ✗ {p}" for p in problems))
            return 1
        print("  ✓ the solution matches the course's calculator to the agora, row by row")
    return 0


if __name__ == "__main__":
    sys.exit(main())
